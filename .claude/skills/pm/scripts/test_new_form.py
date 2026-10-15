"""Tests for new_form.py, the extract_form.py flags and pairing in md_to_docx.py.

Run: uv run --with pytest --with python-docx pytest -q .claude/skills/pm/scripts/test_new_form.py

A blank is generated for every form in references/fields/ into a temporary
directory and each one is rendered to .docx, so a spec that the scaffold or the
renderer cannot handle fails here rather than in front of a teammate.
"""

import contextlib
import io
import json
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import extract_form
import new_form

HERE = Path(__file__).resolve().parent
FIELDS = HERE.parent / "references" / "fields"
FORMS = sorted(p.stem for p in FIELDS.glob("*.json"))
DASHES = re.compile("[\u2013\u2014]")
EMOJI = re.compile("[\U0001F300-\U0001FAFF\u2600-\u26FF\u2700-\u27BF]")

try:
    import md_to_docx
    from docx import Document
except ImportError:  # python-docx missing; the .docx tests skip
    md_to_docx = None


def spec(form):
    return extract_form.load_spec(form)


def blank(form, rows=None):
    return new_form.render(spec(form), rows)


def tables(md):
    """Every pipe table in md as a list of rows of cells, separator rows dropped."""
    out, current = [], []
    for line in md.splitlines():
        if line.startswith("|"):
            if not re.match(r"^\|(\s*---\s*\|)+$", line):
                current.append([c.strip() for c in re.split(r"(?<!\\)\|", line)[1:-1]])
        elif current:
            out.append(current)
            current = []
    if current:
        out.append(current)
    return out


def run_cli(script, *argv):
    return subprocess.run([sys.executable, str(HERE / script), *map(str, argv)],
                          capture_output=True, text=True)


class EveryForm(unittest.TestCase):
    def test_all_66_specs_present(self):
        self.assertEqual(len(FORMS), 66)

    def test_every_blank_renders_with_h1_and_page_headings(self):
        for form in FORMS:
            with self.subTest(form=form):
                s = spec(form)
                md = blank(form)
                self.assertTrue(md.startswith(f"# {s['form']} {s['name']}\n"))
                pages = re.findall(r"^#### .+, page (\d+) of (\d+)$", md, re.M)
                if s["page_count"] > 1:
                    self.assertEqual([int(p) for p, _ in pages], list(range(1, s["page_count"] + 1)))
                    self.assertTrue(all(int(n) == s["page_count"] for _, n in pages))
                else:
                    self.assertEqual(pages, [])

    def test_no_dashes_or_emoji_in_any_blank(self):
        for form in FORMS:
            with self.subTest(form=form):
                md = blank(form)
                self.assertIsNone(DASHES.search(md))
                self.assertIsNone(EMOJI.search(md))

    def test_title_block_uses_printed_labels(self):
        for form in FORMS:
            with self.subTest(form=form):
                s = spec(form)
                md = blank(form)
                self.assertNotIn("**Project:**", md)
                self.assertNotIn("**Date prepared:**", md)
                if not s["title_block"]:
                    self.assertNotRegex(md, r"\*\*(Project Title|Date Prepared|Date):?\*\*")
                    continue
                if form in new_form.TITLE_BLOCK_AS_TABLE:
                    continue
                lines = [f"**{label}:**" for label in s["title_block"]]
                self.assertIn("\n".join(lines) + "\n", md)

    def test_every_table_has_a_separator_and_even_width(self):
        for form in FORMS:
            with self.subTest(form=form):
                for t in tables(blank(form)):
                    self.assertEqual(len({len(r) for r in t}), 1, t[0])

    def test_every_blank_passes_the_fidelity_checker(self):
        try:
            import check_fidelity
        except ImportError:
            self.skipTest("check_fidelity.py not importable")
        if not hasattr(check_fidelity, "check"):
            self.skipTest("check_fidelity.py has no check() in this version")
        for form in FORMS:
            with self.subTest(form=form):
                result = check_fidelity.check(blank(form), check_fidelity.load_spec(form))
                self.assertFalse(result.failed, (result.missing, result.disorder))


class Shapes(unittest.TestCase):
    def test_risk_register_banks_and_grouped_headers(self):
        t = tables(blank("2.32"))
        self.assertEqual(len(t), 2)
        bank1, bank2 = t
        self.assertEqual(bank1[0][:4], ["ID", "Risk Statement", "Owner", "Probability"])
        self.assertEqual(bank1[0][4:8], ["Impact: Scope", "Impact: Quality", "Impact: Schedule", "Impact: Cost"])
        self.assertEqual(bank2[0][0], "Revised Probability")
        self.assertNotIn("ID", bank2[0])
        self.assertEqual(len(bank1), 9)  # header and the eight blank rows each bank prints
        self.assertEqual(len(bank1), len(bank2))

    def test_default_rows_are_the_printed_count(self):
        self.assertEqual(new_form.printed_rows(spec("1.3")), 6)
        self.assertEqual(new_form.printed_rows(spec("3.1")), 7)    # "two banks of seven blank rows"
        self.assertEqual(new_form.printed_rows(spec("1.1")), 7)    # 7 and 6: the larger
        self.assertEqual(new_form.printed_rows(spec("4.2")), 1)    # "1 tall blank row below it"
        self.assertEqual(new_form.printed_rows({"notes": "no count here"}), new_form.DEFAULT_ROWS)
        self.assertEqual(len(tables(blank("1.3"))[0]), 7)

    def test_numbered_lines_are_explained_in_the_reading_paragraph(self):
        paragraph = blank("4.2").split("\n\n")[1]
        self.assertIn("(1) ... (2) ...", paragraph)
        self.assertNotIn("(1)", blank("1.3"))

    def test_dated_copy_path(self):
        import datetime
        path = new_form.default_path(spec("4.2"), Path("/repo"), datetime.date(2026, 10, 5))
        self.assertEqual(path, Path("/repo/docs/forms/4-2-project-status-report-2026-10-05.en.md"))
        self.assertEqual(new_form.default_path(spec("4.2"), Path("/repo")).name, "4-2-project-status-report.en.md")

    def test_rows_option(self):
        bank1 = tables(blank("2.32", rows=7))[0]
        self.assertEqual(len(bank1), 8)
        self.assertTrue(all(c == "" for r in bank1[1:] for c in r))

    def test_group_spanning_every_column_is_a_caption(self):
        md = blank("2.17")
        self.assertIn("**Parametric Estimates**\n\n| ID | Effort Hours |", md)
        self.assertNotIn("Parametric Estimates: ID", md)

    def test_wbs_dictionary_groups_keep_both_totals(self):
        header = tables(blank("2.10"))[1][0]
        self.assertIn("Labor: Total", header)
        self.assertIn("Material: Total", header)

    def test_fields_table(self):
        md = blank("2.10")
        self.assertIn("| Field | Content |\n| --- | --- |\n| **Work Package Name** | |\n"
                      "| **Code of Accounts** | |", md)

    def test_charter_title_block_is_the_first_table(self):
        first = tables(blank("1.1"))[0]
        self.assertEqual(first[0], ["Field", "Content"])
        self.assertEqual([r[0] for r in first[1:]],
                         ["**Project Title**", "**Project Sponsor**", "**Date Prepared**",
                          "**Project Manager**", "**Project Customer**"])

    def test_row_labels_get_a_blank_header_column(self):
        objectives = tables(blank("1.1"))[2]
        self.assertEqual(objectives[0], ["", "Project Objectives", "Success Criteria"])
        self.assertEqual([r[0] for r in objectives[1:]], ["**Scope**", "**Time**", "**Cost**", "**Other**"])

    def test_printed_row_label_column_is_not_doubled(self):
        t = tables(blank("2.37"))[0]
        self.assertEqual(t[0], ["Area", "Integration Approach"])
        self.assertEqual(t[1], ["**Scope**", ""])
        process = tables(blank("4.5"))[2]
        self.assertEqual(process[0], ["Process", "Followed", "Tools and Techniques Used"])
        plans = tables(blank("2.1"))[3]
        self.assertEqual(plans[0], ["Name", "Comment"])
        self.assertEqual(plans[1], ["**Scope**", ""])
        matrix = tables(blank("2.25"))[0]
        self.assertEqual(matrix[0][:2], ["", "Person 1"])
        self.assertEqual(len(matrix[0]), 6)

    def test_signatures(self):
        sig = tables(blank("1.1"))[-1]
        self.assertEqual(sig[0], ["", "Project Manager", "Sponsor or Originator"])
        self.assertEqual([r[0] for r in sig[1:]], ["**Signature**", "**Name**", "**Date**"])
        team = tables(blank("2.27", rows=4))[-1]
        self.assertEqual(team[0], ["Signature", "Date"])
        self.assertEqual(len(team), 5)

    def test_checkboxes_are_task_list_items(self):
        md = blank("3.6")
        self.assertIn("**Area Audited**\n\n- [ ] Project processes\n- [ ] Project documents\n", md)
        self.assertNotIn("[x]", md)

    def test_picture_gets_table_and_mermaid(self):
        md = blank("2.18")
        self.assertEqual(md.count("```mermaid\ngantt"), 2)
        self.assertIn("| ID | WBS | Task Name | Start | Finish |", md)
        self.assertNotIn("| Phase |", md)
        cost = blank("2.22")
        self.assertIn("| Period | Cumulative Planned Value |", cost)
        self.assertIn("```mermaid\nxychart-beta", cost)
        net = blank("2.15")
        self.assertIn("```mermaid\nflowchart LR", net)
        self.assertIn("| Activity | Predecessor | Relationship | Lead or Lag |", net)

    def test_picture_heading_kept_and_sample_caption_dropped(self):
        roadmap = blank("2.3")
        self.assertIn("**Life Cycle Phases**", roadmap)
        self.assertIn("**Timeline**", roadmap)
        self.assertNotIn("Sample Gantt Chart", blank("2.18"))

    def test_sample_legend_dropped_and_key_legend_kept(self):
        self.assertNotIn("two-day lead", blank("2.15"))
        self.assertIn("R = Responsible: The person performing the work.", blank("2.25"))

    def test_no_sample_content(self):
        self.assertNotIn("Kitchen", blank("2.18"))
        self.assertNotIn("Origination", blank("2.3"))

    def test_narrative_is_verbatim(self):
        md = blank("6.1")
        for block in spec("6.1")["pages"][0]["blocks"]:
            self.assertIn("\n" + block["text"] + "\n", md)
        self.assertNotIn("| Field |", md)

    def test_outline_is_a_nested_ordered_list(self):
        md = blank("2.9")
        self.assertIn("1. Project\n   1. Major Deliverable\n      1. Control Account\n"
                      "         1. Work Package\n         2. Work Package\n", md)
        self.assertIn("\n      2. Work Package\n   2. Control Account\n", md)
        rbs = blank("2.29")
        self.assertIn("\n   5. Locations\n      1. Location 1\n", rbs)

    def test_text_heading_bold_and_instruction_plain(self):
        md = blank("2.2")
        self.assertIn("\nAttach relevant forms used in the change control process.\n", md)
        self.assertIn("**Definitions of Change**\n\n| Field | Content |\n| --- | --- |\n"
                      "| **Schedule change** | |", md)
        self.assertNotIn("| **Change Control Board** |", md)
        self.assertIn("**Change Control Process**\n\n| Field | Content |\n| --- | --- |\n"
                      "| **Change request submittal** | |", md)
        charter = blank("1.1")
        self.assertIn("\n**Project Manager Authority Level**\n", charter)
        self.assertIn("\n**Approvals**\n", charter)
        self.assertIn("\n**Baselines**\n\nAttach all project baselines.\n", blank("2.1"))
        self.assertIn("**Roles and Responsibilities**", blank("2.31"))

    def test_asterisk_in_a_label_is_escaped(self):
        self.assertIn(r"**EAC w/ CPI\*SPI [AC + ((BAC - EV)/(CPI\*SPI))]**", blank("4.4"))

    def test_title_block_repeats_only_where_printed(self):
        rmp = blank("2.31")
        self.assertEqual(rmp.count("**Project Title:**"), 2)
        self.assertEqual(new_form.title_block_pages(spec("2.31")), {1, 2})
        self.assertEqual(new_form.title_block_pages(spec("2.7")), {1, 2})
        self.assertEqual(new_form.title_block_pages(spec("2.33")), {1})
        self.assertEqual(blank("2.33").count("**Project Title:**"), 1)

    def test_page_with_its_own_title_bar(self):
        self.assertIn("#### INTER-REQUIREMENTS TRACEABILITY MATRIX, page 2 of 2", blank("2.7"))
        self.assertIn("#### BOTTOM-UP COST ESTIMATING WORKSHEET, page 2 of 2", blank("2.21"))
        self.assertIn("#### REQUIREMENTS TRACEABILITY MATRIX, page 1 of 2", blank("2.7"))

    def test_odd_title_blocks(self):
        self.assertIn("**Project Name:**\n\n", blank("6.3"))
        self.assertNotIn("**Date Prepared:**", blank("6.3"))
        self.assertIn("**Project Title:**\n**Date:**\n", blank("2.11"))
        self.assertIn("**Project Auditor:**\n**Audit Date:**", blank("3.6"))

    def test_citation(self):
        self.assertIn("*Form 2.32, A Project Manager's Book of Forms, 3rd edition, "
                      "PDF page 141 (printed page 130).", blank("2.32"))
        self.assertIn("PDF pages 16 to 19 (printed pages 5 to 8).*", blank("1.1"))
        self.assertNotIn("**Source:**", blank("1.1"))


class Cli(unittest.TestCase):
    def test_default_path(self):
        root = Path("/repo")
        self.assertEqual(new_form.default_path(spec("2.32"), root),
                         root / "docs" / "forms" / "2-32-risk-register.en.md")
        self.assertEqual(new_form.default_path(spec("4.9"), root).name, "4-9-product-acceptance-form.en.md")

    def test_writes_refuses_overwrite_and_forces(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "sub" / "rr.md"
            first = run_cli("new_form.py", "risk register", "--out", out)
            self.assertEqual(first.returncode, 0, first.stderr)
            self.assertEqual(out.read_text(encoding="utf-8"), blank("2.32"))
            out.write_text("edited", encoding="utf-8")
            again = run_cli("new_form.py", "2.32", "--out", out)
            self.assertNotEqual(again.returncode, 0)
            self.assertIn("--force", again.stderr)
            self.assertEqual(out.read_text(encoding="utf-8"), "edited")
            forced = run_cli("new_form.py", "2.32", "--out", out, "--force", "--rows", "5")
            self.assertEqual(forced.returncode, 0, forced.stderr)
            self.assertEqual(out.read_text(encoding="utf-8"), blank("2.32", rows=5))

    def test_dated_copy_into_a_scratch_directory_never_touches_docs(self):
        with tempfile.TemporaryDirectory() as tmp:
            res = run_cli("new_form.py", "4.2", "--out", tmp + "/", "--dated", "2026-10-19")
            self.assertEqual(res.returncode, 0, res.stderr)
            files = sorted(p.name for p in Path(tmp).iterdir())
            self.assertEqual(files, ["4-2-project-status-report-2026-10-19.en.md"])
            self.assertEqual((Path(tmp) / files[0]).read_text(encoding="utf-8"), blank("4.2"))
            plain = run_cli("new_form.py", "4.2", "--out", tmp)
            self.assertEqual(plain.returncode, 0, plain.stderr)
            self.assertIn("4-2-project-status-report.en.md", sorted(p.name for p in Path(tmp).iterdir()))
            bad = run_cli("new_form.py", "4.2", "--out", tmp, "--dated", "19 October")
            self.assertNotEqual(bad.returncode, 0)

    def test_status_report_reading_paragraph_says_schedule_first_then_cost(self):
        text = blank("4.2")
        italic = text.splitlines()[2]
        self.assertIn("the first copy of each is the schedule one and the second the cost one", italic)
        self.assertIn("PDF pages 207 to 208 (printed pages 196 to 197)", italic)
        self.assertIn("PDF pages 209 to 212 (printed pages 198 to 201)", italic)
        self.assertNotIn("schedule one", blank("2.32").splitlines()[2])

    def test_status_report_risk_and_issue_tables_print_one_blank_row(self):
        text = blank("4.2")
        for header in ("| Risk |", "| Issue |"):
            body = text.split(header + "\n| --- |\n")[1].split("\n\n")[0]
            self.assertEqual(body.splitlines(), ["| |"], header)

    def test_no_citation_is_nested_inside_another_bracket(self):
        italic = blank("4.2").splitlines()[2]
        self.assertNotIn("(Book of Forms form 4.2", italic)
        self.assertIn("Table 4.2, PDF pages 207 to 208 (printed pages 196 to 197)", italic)

    def test_no_out_prints_and_writes_nothing(self):
        before = sorted(p.name for p in (HERE.parents[3] / "docs" / "forms").glob("*"))
        res = run_cli("new_form.py", "6.4")
        self.assertEqual(res.returncode, 0)
        self.assertEqual(res.stdout, blank("6.4"))
        self.assertEqual(before, sorted(p.name for p in (HERE.parents[3] / "docs" / "forms").glob("*")))

    def test_stdout(self):
        res = run_cli("new_form.py", "6.4", "--out", "-")
        self.assertEqual(res.returncode, 0, res.stderr)
        self.assertEqual(res.stdout, blank("6.4"))

    def test_bad_rows(self):
        self.assertNotEqual(run_cli("new_form.py", "6.4", "--out", "-", "--rows", "0").returncode, 0)


FAKE_PROSE = "\n".join([
    "2.32 RISK REGISTER", "Purpose text.",
    "Tailoring Tips", "  Tip one.",
    "Alignment", "  Risk management plan",
    "Description", "Element   Description", "ID   A unique identifier",
])


def run_extract(*argv):
    """Printed sections of extract_form.main with the PDF stubbed out."""
    buf = io.StringIO()
    with mock.patch.object(extract_form, "find_pdf", return_value=Path("book.pdf")), \
            mock.patch.object(extract_form, "page_text",
                              side_effect=lambda pdf, a, b: FAKE_PROSE if a < 141 else "BLANK PAGE TEXT"), \
            mock.patch.object(sys, "argv", ["extract_form.py", *argv]), \
            contextlib.redirect_stdout(buf):
        extract_form.main()
    out = buf.getvalue()
    return out, re.findall(r"^## (.+)$", out, re.M)


class ExtractFlags(unittest.TestCase):
    def test_default(self):
        _, sections = run_extract("2.32")
        self.assertEqual(sections, ["Document elements", "Printed form"])

    def test_alignment_and_tailoring_add_to_default(self):
        out, sections = run_extract("2.32", "--alignment", "--tailoring")
        self.assertEqual(sections, ["Tailoring tips", "Alignment", "Document elements", "Printed form"])
        self.assertIn("Risk management plan", out)
        self.assertIn("BLANK PAGE TEXT", out)

    def test_explicit_selector_still_selects(self):
        _, sections = run_extract("2.32", "--blank", "--alignment")
        self.assertEqual(sections, ["Alignment", "Printed form"])

    def test_fields_needs_no_pdf(self):
        buf = io.StringIO()
        with mock.patch.object(extract_form, "find_pdf", side_effect=AssertionError("pdf read")), \
                mock.patch.object(sys, "argv", ["extract_form.py", "2.32", "--fields"]), \
                contextlib.redirect_stdout(buf):
            extract_form.main()
        out = buf.getvalue()
        self.assertIn("## Field spec, references/fields/2.32.json", out)
        self.assertIn("Impact: Scope", out)
        self.assertIn("table, bank 2: Revised Probability", out)

    def test_fields_with_alignment(self):
        _, sections = run_extract("2.32", "--fields", "--alignment")
        self.assertEqual(sections, ["Field spec, references/fields/2.32.json", "Alignment"])

    def test_describe_every_spec(self):
        for form in FORMS:
            with self.subTest(form=form):
                text = extract_form.describe_spec(spec(form))
                self.assertTrue(text.startswith("Title block: "))
                self.assertEqual(text.count("\nPage "), spec(form)["page_count"])

    def test_flat_columns(self):
        cols = ["ID", {"group": "Impact", "columns": ["Scope", "Cost"]}, "Score"]
        self.assertEqual(extract_form.flat_columns(cols), (["ID", "Impact: Scope", "Impact: Cost", "Score"], ""))
        caption = [{"group": "Parametric Estimates", "columns": ["ID", "Cost"]}]
        self.assertEqual(extract_form.flat_columns(caption), (["ID", "Cost"], "Parametric Estimates"))


def docx_rows(path):
    """Each table row of a .docx as its distinct cell texts, merged cells once."""
    rows = []
    for table in Document(path).tables:
        for row in table.rows:
            seen, cells = [], []
            for c in row.cells:
                if c._tc not in seen:
                    seen.append(c._tc)
                    cells.append(c.text)
            rows.append(cells)
    return rows


def docx_paragraphs(path):
    return [p.text for p in Document(path).paragraphs]


@unittest.skipIf(md_to_docx is None, "python-docx not installed")
class Docx(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.tmp.name)
        for form in FORMS:
            md = cls.dir / f"{form}.md"
            md.write_text(blank(form), encoding="utf-8")
            md_to_docx.render(md, cls.dir / f"{form}.docx")

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def path(self, form):
        return self.dir / f"{form}.docx"

    def test_every_form_rendered(self):
        for form in FORMS:
            self.assertTrue(self.path(form).stat().st_size > 0, form)

    def test_form_number_from_h1(self):
        self.assertEqual(md_to_docx.form_number("# 2.10 WBS Dictionary\n\ntext"), "2.10")
        self.assertIsNone(md_to_docx.form_number("# Project Charter Package\n"))
        self.assertIsNone(md_to_docx.load_spec(None))
        self.assertIsNone(md_to_docx.load_spec("9.99"))

    def test_fields_pairs_share_a_row(self):
        rows = docx_rows(self.path("2.10"))
        self.assertIn(["Work Package Name", "", "Code of Accounts", ""], rows)
        self.assertIn(["Description of Work", "", "Assumptions and Constraints", ""], rows)
        self.assertIn(["Milestones", "", "Due Dates", ""], rows)
        self.assertIn(["Quality Requirements", ""], rows)

    def test_triple_pair(self):
        rows = docx_rows(self.path("2.11"))
        self.assertIn(["Level of Accuracy", "", "Units of Measure", "", "Variance Thresholds", ""], rows)
        self.assertIn(["Schedule Methodology", ""], rows)

    def test_title_block_pairs(self):
        rows = docx_rows(self.path("3.6"))
        self.assertIn(["Project Title:", "", "Date Prepared:", ""], rows)
        self.assertIn(["Project Auditor:", "", "Audit Date:", ""], rows)
        charter = docx_rows(self.path("1.1"))
        self.assertIn(["Project Title", ""], charter)
        self.assertIn(["Project Sponsor", "", "Date Prepared", ""], charter)
        self.assertIn(["Project Manager", "", "Project Customer", ""], charter)

    def test_no_pairs_no_change(self):
        paras = docx_paragraphs(self.path("6.3"))
        self.assertIn("Project Name:", paras)
        self.assertFalse(any(r and r[0] == "Project Name:" for r in docx_rows(self.path("6.3"))))

    def test_field_content_header_not_printed(self):
        for form in ("2.2", "2.10", "1.1"):
            self.assertNotIn(["Field", "Content"], docx_rows(self.path(form)), form)

    def test_title_pair_from_the_later_forms(self):
        self.assertIn(["Project Title:", "", "Date Prepared:", ""], docx_rows(self.path("2.32")))
        self.assertIn(["Project Title:", "", "Date:", ""], docx_rows(self.path("2.37")))

    def test_pair_only_in_printed_order(self):
        import gdoc_edit  # on sys.path once render() has run
        rows = [["**Code of Accounts**", "x"], ["**Work Package Name**", "y"]]
        groups = md_to_docx.pair_rows(rows, [["Work Package Name", "Code of Accounts"]], gdoc_edit.parse_inline)
        self.assertEqual([len(g) for g in groups], [1, 1])

    def test_printed_title_bar(self):
        first = docx_rows(self.path("4.9"))[0]
        self.assertEqual(first, ["FORMAL ACCEPTANCE FORM"])
        self.assertIn(["INTER-REQUIREMENTS TRACEABILITY MATRIX"], docx_rows(self.path("2.7")))

    def test_multi_page_form_has_one_bar_per_page(self):
        bars = [r for r in docx_rows(self.path("3.3")) if r == ["CHANGE REQUEST"]]
        self.assertEqual(len(bars), 3)

    def test_mermaid_skipped(self):
        text = "\n".join(docx_paragraphs(self.path("2.18")))
        self.assertNotIn("gantt", text)
        self.assertNotIn("mermaid", text)
        self.assertIn(["ID", "WBS", "Task Name", "Start", "Finish"], docx_rows(self.path("2.18")))

    def test_outline_keeps_codes_and_depth(self):
        doc = Document(self.path("2.9"))
        paras = {p.text: p.paragraph_format.left_indent for p in doc.paragraphs}
        self.assertIn("1.1.1.1. Work Package", paras)
        self.assertGreater(paras["1.1.1.1. Work Package"], paras["1.1. Major Deliverable"])
        self.assertIn("1.3.2.2. Work Package", paras)

    def test_checkboxes_render_as_boxes(self):
        paras = docx_paragraphs(self.path("3.6"))
        self.assertIn("\u2610 Project processes", paras)

    def test_escaped_asterisk_survives(self):
        flat = [c for r in docx_rows(self.path("4.4")) for c in r]
        self.assertIn("EAC w/ CPI*SPI [AC + ((BAC - EV)/(CPI*SPI))]", flat)

    def test_form_option_pairs_an_unnumbered_document(self):
        md = self.dir / "charter.md"
        md.write_text("# Project Charter Package\n\n#### PROJECT CHARTER, page 1 of 4\n\n"
                      "| Field | Content |\n| --- | --- |\n| **Project Title** | LCMS |\n"
                      "| **Project Sponsor** | S |\n| **Date Prepared** | 2 September 2026 |\n",
                      encoding="utf-8")
        plain = self.dir / "plain.docx"
        md_to_docx.render(md, plain)
        self.assertNotIn(["Project Sponsor", "S", "Date Prepared", "2 September 2026"], docx_rows(plain))
        paired = self.dir / "paired.docx"
        md_to_docx.render(md, paired, form="1.1")
        self.assertIn(["Project Sponsor", "S", "Date Prepared", "2 September 2026"], docx_rows(paired))

    def test_flat_numbered_list_untouched(self):
        md = "1. one\n2. two\n\n1. a\n   1. b\n"
        out = md_to_docx.nest_ordered_lists(md)
        self.assertTrue(out.startswith("1. one\n2. two\n"))
        self.assertIn("```\n1. a\n  1.1. b\n```", out)


class SpecShape(unittest.TestCase):
    """The spec features the scaffold relies on, so a malformed spec fails here."""

    def test_pairs_are_adjacent_labels_in_printed_order(self):
        for form in FORMS:
            s = spec(form)
            runs = [s["title_block"]] + [b["labels"] for p in s["pages"] for b in p["blocks"]
                                         if b["type"] == "fields"]
            for pair in s["pairs"]:
                with self.subTest(form=form, pair=pair):
                    self.assertTrue(any(
                        run[i:i + len(pair)] == pair for run in runs for i in range(len(run))), pair)

    def test_block_types_known(self):
        known = {"fields", "table", "text", "checkboxes", "signatures", "picture"}
        for form in FORMS:
            s = spec(form)
            self.assertEqual(len(s["pages"]), s["page_count"], form)
            for p in s["pages"]:
                for b in p["blocks"]:
                    self.assertIn(b["type"], known, form)

    def test_title_block_printed_on_one_line_is_paired(self):
        for form in FORMS:
            s = spec(form)
            if len(s["title_block"]) == 2:
                self.assertIn(s["title_block"], s["pairs"], form)

    def test_page_and_table_keys(self):
        for form in FORMS:
            for p in spec(form)["pages"]:
                self.assertLessEqual(set(p), {"page", "pdf_page", "blocks", "title_block", "printed_title"}, form)
                for b in p["blocks"]:
                    if b.get("row_label_column"):
                        self.assertTrue(b["row_labels"] and isinstance(b["columns"][0], str), form)

    def test_spec_files_are_valid_json(self):
        for path in FIELDS.glob("*.json"):
            json.loads(path.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
